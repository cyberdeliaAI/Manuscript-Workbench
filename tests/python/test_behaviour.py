"""Explicit, readable assertions about behaviour that must keep working (complements the golden snapshot)."""
from __future__ import annotations

import io
import json

import pytest

from snapshot import rtf


# ---------- text decoding / import ----------

def test_decode_windows_1252_and_crlf(mw):
    text, notes = mw.decode_text_bytes("Café “quoted”\r\nnext".encode("cp1252"))
    assert text == "Café “quoted”\nnext"
    assert "Read as Windows-1252" in notes and "CRLF → LF" in notes


def test_decode_utf8_bom_is_removed(mw):
    text, notes = mw.decode_text_bytes(b"\xef\xbb\xbfhello")
    assert text == "hello" and "Read as UTF-8 BOM" in notes


def test_rtf_to_text_handles_escapes(mw):
    text = mw.rtf_to_text(rtf("{@rtf1@ansi Caf@'e9 @u8220?quoted@u8221?@par Next@tab line@par}"))
    assert text.startswith("Café “quoted”")
    assert "Next\tline" in text


def test_import_old_doc_is_refused(mw):
    with pytest.raises(ValueError, match=r"\.doc"):
        mw.import_file("book.doc", b"x")


def test_import_detects_content_by_signature(mw):
    result = mw.import_file("noext", rb"{\rtf1\ansi Hello\par}")
    assert result["text"] == "Hello"
    assert any("Detected RTF" in n for n in result["notes"])


# ---------- markdown preview / sanitizer ----------

def test_preview_scene_break_and_blocks(mw):
    html = mw.render_markdown_html("Text\n\n***\n\n::: verse\nLine one\n:::\n")
    assert 'class="scene"' in html
    assert "block-verse" in html


@pytest.mark.parametrize("vector", [
    "<script>alert(1)</script>",
    '<img src="x" onerror="alert(1)">',
    '<a href="javascript:alert(1)">x</a>',
    '<a href=" javascript:alert(1)">x</a>',
    '<iframe src="http://x"></iframe>',
    '<div style="position:fixed">x</div>',
    '<svg onload="alert(1)"></svg>',
])
def test_sanitizer_blocks_active_content(mw, vector):
    out = mw.sanitize_markdown_html(vector).lower()
    for bad in ("<script", "onerror", "javascript:", "<iframe", "style=", "onload", "<svg"):
        assert bad not in out


# ---------- preflight ----------

def codes(result):
    return {issue["code"] for issue in result["issues"]}


def test_preflight_flags_missing_chapters_and_candidates(mw, corpus):
    result = mw.preflight(corpus["chapter_candidates"], {})
    assert {"no-chapters", "chapter-candidates", "missing-title", "missing-author"} <= codes(result)
    assert result["summary"]["error"] == 1


def test_preflight_broken_structure(mw, corpus):
    result = mw.preflight(corpus["structure_broken"], {"title": "T", "author": "A"})
    assert "unclosed-structure-block" in codes(result)
    assert "chapter-number-sequence" in codes(result)


def test_preflight_messy_pdf_text(mw, corpus):
    result = mw.preflight(corpus["messy_pdf_en"], {})
    found = codes(result)
    assert {"page-numbers", "repeated-lines", "invisible-chars", "control-chars", "mojibake", "tabs",
            "extra-blank-runs", "toc-detected"} <= found


def test_preflight_warns_that_lists_and_tables_stay_literal_in_docx(mw, corpus):
    found = codes(mw.preflight(corpus["structure_en"], {"title": "T", "author": "A"}))
    assert {"docx-list-literal", "docx-table-literal"} <= found


# ---------- DOCX export / import ----------

def test_docx_export_maps_structure_to_styles(mw, corpus):
    from docx import Document

    data = mw.create_manuscript_docx(corpus["structure_en"], {"title": "T", "author": "A"})
    doc = Document(io.BytesIO(data))
    styles = {p.text: p.style.name for p in doc.paragraphs}
    assert styles["Part One"] == "Heading 1"
    assert styles["Chapter One"] == "Heading 1"
    assert styles["A Subhead"] == "Heading 2"
    assert styles["Lower subhead"] == "Heading 3"
    assert styles["A quoted paragraph."] == "Manuscript Block Quote"
    assert styles["Roses are red"] == "Manuscript Verse"
    assert styles["A note written by a character."] == "Manuscript Written Note"
    assert styles["* * *"] == "Normal"
    assert doc.core_properties.title == "T" and doc.core_properties.author == "A"


def test_docx_inline_markers_become_run_formatting(mw):
    from docx import Document

    data = mw.create_manuscript_docx("a *it* **bo** ***bi*** [[u:un]] [[sc:sc]] `mo`", {})
    runs = {r.text: r for r in Document(io.BytesIO(data)).paragraphs[0].runs}
    assert runs["it"].italic and runs["bo"].bold and runs["bi"].bold and runs["bi"].italic
    assert runs["un"].underline
    assert runs["sc"].style.name == "Manuscript Small Caps"
    assert runs["mo"].style.name == "Manuscript Monospace"


def test_docx_roundtrip_restores_markdown_structure(mw):
    source = "# Chapter\n\n## Sub\n\n> Quote\n\n::: verse\nLine\n:::\n\nPlain *it* and **bo**."
    text, _notes = mw.import_docx_bytes(mw.create_manuscript_docx(source, {}))
    assert "# Chapter" in text and "## Sub" in text and "> Quote" in text
    assert "::: verse\nLine\n:::" in text
    assert "*it*" in text and "**bo**" in text


# ---------- provider helpers ----------

def test_normalize_base_url(mw):
    assert mw.normalize_base_url("", "lmstudio") == "http://localhost:1234/v1"
    assert mw.normalize_base_url("localhost:1234", "lmstudio") == "http://localhost:1234/v1"
    assert mw.normalize_base_url("http://x/v1/", "custom") == "http://x/v1"
    assert mw.normalize_base_url("", "omlx") == "http://localhost:8000/v1"


# ---------- HTTP surface (legacy protocol; the stream format changes in phase 1) ----------

def test_http_index_health_and_404(server):
    status, headers, body = server.request("GET", "/")
    assert status == 200 and b"Manuscript Workbench" in body
    status, _h, body = server.request("GET", "/api/health")
    assert status == 200 and json.loads(body)["ok"] is True
    assert server.request("GET", "/nope")[0] == 404
    assert server.request("POST", "/nope", body={})[0] == 404


def test_http_help_page_renders_markdown(server):
    status, _h, body = server.request("GET", "/help")
    assert status == 200 and b"Manuscript Workbench" in body


def test_http_render_preflight_and_export_endpoints(server):
    status, _h, body = server.request("POST", "/api/render_markdown", body={"text": "# T\n\ntext"})
    assert status == 200 and "<h1>" in json.loads(body)["html"]
    status, _h, body = server.request("POST", "/api/preflight", body={"text": "# T", "metadata": {}})
    assert status == 200 and "issues" in json.loads(body)
    status, headers, body = server.request("POST", "/api/export_docx", body={"text": "# T", "metadata": {}, "filename": "my book"})
    assert status == 200 and body[:2] == b"PK"
    assert "my-book-formatted.docx" in headers["Content-Disposition"]


def test_http_import_endpoint_roundtrip(server):
    import base64

    payload = {"filename": "a.md", "data_b64": base64.b64encode(b"# Title\r\ntext").decode()}
    status, _h, body = server.request("POST", "/api/import", body=payload)
    assert status == 200 and json.loads(body)["text"] == "# Title\ntext"


def test_http_settings_roundtrip(server):
    status, _h, _b = server.request("POST", "/api/settings", body={"settings": {"wrap": True, "lmProvider": "omlx"}})
    assert status == 200
    status, _h, body = server.request("GET", "/api/settings")
    data = json.loads(body)
    assert data["settings"]["wrap"] is True and data["settings"]["lmProvider"] == "omlx"
    assert server.settings_path.exists()


def test_http_lm_models_and_auth(server, mock_provider):
    status, _h, body = server.request("POST", "/api/lm_models", body={"base_url": mock_provider.base_url, "provider": "custom"})
    assert status == 200 and json.loads(body)["models"] == ["normal", "think", "reasoning"]


def test_http_lm_models_with_api_key(server):
    from mock_openai import MockOpenAI

    provider = MockOpenAI(api_key="secret").start()
    try:
        ok = server.request("POST", "/api/lm_models", body={"base_url": provider.base_url, "api_key": "secret"})
        assert ok[0] == 200
        bad = server.request("POST", "/api/lm_models", body={"base_url": provider.base_url, "api_key": "wrong"})
        assert bad[0] == 502 and "401" in json.loads(bad[2])["error"] and "API Key" in json.loads(bad[2])["error"]
    finally:
        provider.stop()
