"""DOCX import in document order: tables, lists, hyperlinks, and what is still skipped."""
from __future__ import annotations

import io
import re
import struct
import zipfile
import zlib

import pytest

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def saved(doc) -> bytes:
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def new_doc():
    from docx import Document

    return Document()


def imported(mw, doc):
    text, notes = mw.import_docx_bytes(saved(doc))
    return text, notes


def png_bytes() -> bytes:
    def chunk(kind: bytes, data: bytes) -> bytes:
        return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data))

    raw = b"\x00\xff\x00\x00"
    return b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0)) + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b"")


def test_a_table_comes_back_in_its_place_with_blank_lines_around_it(mw):
    doc = new_doc()
    doc.add_paragraph("Before.")
    table = doc.add_table(rows=2, cols=2)
    for r, row in enumerate([["A", "B"], ["1", "2"]]):
        for c, text in enumerate(row):
            table.cell(r, c).text = text
    doc.add_paragraph("After.")
    text, notes = imported(mw, doc)
    assert text == "Before.\n\n| A | B |\n| --- | --- |\n| 1 | 2 |\n\nAfter."
    assert "Converted DOCX tables to Markdown tables: 1" in notes
    assert not any("Skipped DOCX tables" in note for note in notes)


def test_table_cells_keep_inline_formatting_but_the_header_is_not_marked_bold(mw):
    doc = new_doc()
    table = doc.add_table(rows=2, cols=2)
    head = table.cell(0, 0).paragraphs[0].add_run("Head")
    head.bold = True
    table.cell(0, 1).text = "Plain"
    table.cell(1, 0).paragraphs[0].add_run("bold").bold = True
    table.cell(1, 1).text = "a | b"
    text, _ = imported(mw, doc)
    assert text.splitlines() == ["| Head | Plain |", "| --- | --- |", "| **bold** | a \\| b |"]


def test_alignment_and_multi_paragraph_cells(mw):
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = new_doc()
    table = doc.add_table(rows=2, cols=3)
    for c, name in enumerate(["L", "C", "R"]):
        table.cell(0, c).text = name
    table.cell(1, 0).text = "x"
    table.cell(1, 0).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
    table.cell(1, 1).text = "y"
    table.cell(1, 1).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    table.cell(1, 2).text = "first"
    table.cell(1, 2).add_paragraph("second")
    table.cell(1, 2).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    text, _ = imported(mw, doc)
    assert text.splitlines()[1] == "| :--- | :---: | ---: |"
    assert text.splitlines()[2] == "| x | y | first<br>second |"


def test_merged_cells_do_not_repeat_their_text(mw):
    doc = new_doc()
    table = doc.add_table(rows=2, cols=3)
    table.cell(0, 0).text = "H1"
    table.cell(0, 1).text = "H2"
    table.cell(0, 2).text = "H3"
    merged = table.cell(1, 0).merge(table.cell(1, 1))
    merged.text = "wide"
    table.cell(1, 2).text = "end"
    text, _ = imported(mw, doc)
    assert text.splitlines()[2] == "| wide |  | end |"


def test_a_one_column_table_is_still_a_table(mw):
    doc = new_doc()
    table = doc.add_table(rows=2, cols=1)
    table.cell(0, 0).text = "Box"
    table.cell(1, 0).text = "content"
    text, _ = imported(mw, doc)
    assert text == "| Box |\n| --- |\n| content |"
    assert "<table>" in mw.render_markdown_html(text)


def test_word_list_styles_become_markdown_lists(mw):
    doc = new_doc()
    doc.add_paragraph("Intro.")
    for style, text in [("List Bullet", "one"), ("List Bullet", "two"), ("List Bullet 2", "nested"), ("List Bullet", "three")]:
        doc.add_paragraph(text, style=style)
    for text in ("first", "second", "third"):
        doc.add_paragraph(text, style="List Number")
    doc.add_paragraph("Outro.")
    text, notes = imported(mw, doc)
    assert text == "Intro.\n\n- one\n- two\n    - nested\n- three\n1. first\n2. second\n3. third\n\nOutro."
    assert "Converted DOCX list items to Markdown lists: 7" in notes


def test_lists_built_from_numbering_definitions_use_their_format_and_levels(mw):
    from docx.oxml import parse_xml

    doc = new_doc()
    root = doc.part.numbering_part.element
    abstract = parse_xml(
        f'<w:abstractNum xmlns:w="{W}" w:abstractNumId="90"><w:multiLevelType w:val="multilevel"/>'
        '<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:lvlText w:val="%1."/></w:lvl>'
        '<w:lvl w:ilvl="1"><w:start w:val="1"/><w:numFmt w:val="bullet"/><w:lvlText w:val="o"/></w:lvl></w:abstractNum>'
    )
    root.insert(0, abstract)
    root.append(parse_xml(f'<w:num xmlns:w="{W}" w:numId="90"><w:abstractNumId w:val="90"/></w:num>'))
    for text, ilvl in [("step one", 0), ("a detail", 1), ("step two", 0)]:
        p = doc.add_paragraph(text, style="List Paragraph")
        numbers = p._p.get_or_add_pPr().get_or_add_numPr()
        numbers.get_or_add_ilvl().val = ilvl
        numbers.get_or_add_numId().val = 90
    text, _ = imported(mw, doc)
    assert text == "1. step one\n    - a detail\n2. step two"


def test_numbered_headings_stay_headings(mw):
    doc = new_doc()
    heading = doc.add_heading("Chapter", level=1)
    numbers = heading._p.get_or_add_pPr().get_or_add_numPr()
    numbers.get_or_add_ilvl().val = 0
    numbers.get_or_add_numId().val = 1
    text, _ = imported(mw, doc)
    assert text == "# Chapter"


def test_a_list_paragraph_without_text_is_not_a_list_item(mw):
    doc = new_doc()
    doc.add_paragraph("", style="List Bullet")
    doc.add_paragraph("real", style="List Bullet")
    text, _ = imported(mw, doc)
    assert text == "\n- real"


def test_hyperlink_text_is_kept_and_external_links_become_markdown_links(mw):
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml import parse_xml

    doc = new_doc()
    p = doc.add_paragraph("See ")
    rid = doc.part.relate_to("https://example.org/page", RT.HYPERLINK, is_external=True)
    p._p.append(parse_xml(
        f'<w:hyperlink xmlns:w="{W}" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" r:id="{rid}"><w:r><w:t>the page</w:t></w:r></w:hyperlink>'
    ))
    p.add_run(" and ")
    p._p.append(parse_xml(f'<w:hyperlink xmlns:w="{W}" w:anchor="somewhere"><w:r><w:t>a bookmark</w:t></w:r></w:hyperlink>'))
    text, _ = imported(mw, doc)
    assert text == "See [the page](https://example.org/page) and a bookmark"


def test_images_are_still_skipped_but_now_reported(mw, tmp_path):
    image = tmp_path / "dot.png"
    image.write_bytes(png_bytes())
    doc = new_doc()
    doc.add_paragraph("Text")
    doc.add_picture(str(image))
    _text, notes = imported(mw, doc)
    assert "Skipped DOCX images: 1" in notes


def test_round_trip_of_lists_and_tables(mw):
    source = "\n".join([
        "# Title", "", "- one", "- two *it*", "    - nested", "- three", "", "5. five", "6. six", "",
        "| Name | Qty |", "| --- | ---: |", "| Apple | 3 |", "| Pear | **10** |", "", "After.",
    ])
    text, _ = mw.import_docx_bytes(mw.create_manuscript_docx(source, {}))
    assert text.strip() == source


def test_round_trip_with_lists_switched_off_keeps_literal_markers(mw):
    text, notes = mw.import_docx_bytes(mw.create_manuscript_docx("- a\n- b", {}, {"lists": False}))
    assert text.strip() == "- a\n- b" and not any("list items" in n for n in notes)


def test_markdown_links_become_real_word_hyperlinks_and_come_back(mw):
    source = "See [the page](https://example.org/a?x=1&y=2), **bold** and [mail](mailto:a@b.c); [not a link](ftp://x) stays text."
    data = mw.create_manuscript_docx(source, {})
    from docx import Document
    from docx.opc.constants import RELATIONSHIP_TYPE as RT

    doc = Document(io.BytesIO(data))
    targets = sorted(r.target_ref for r in doc.part.rels.values() if r.reltype == RT.HYPERLINK)
    assert targets == ["https://example.org/a?x=1&y=2", "mailto:a@b.c"]
    assert doc.paragraphs[0].text == "See the page, bold and mail; [not a link](ftp://x) stays text."
    text, _ = mw.import_docx_bytes(data)
    assert text.strip() == source


def test_links_work_inside_list_items_and_table_cells(mw):
    source = "- item with [a link](https://a.example)\n\n| h | k |\n|---|---|\n| [cell](https://c.example) | x |"
    text, _ = mw.import_docx_bytes(mw.create_manuscript_docx(source, {}))
    assert "- item with [a link](https://a.example)" in text and "| [cell](https://c.example) | x |" in text


def test_markdown_images_are_not_mistaken_for_links(mw):
    source = "A [link](https://example.com) and ![image](https://example.com/x.png)."
    from docx import Document

    doc = Document(io.BytesIO(mw.create_manuscript_docx(source, {})))
    assert doc.paragraphs[0].text == "A link and ![image](https://example.com/x.png)."


# ---- the whole corpus through DOCX and back -----------------------------------------------------------

CONTROL = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f]")


def words(text: str):
    """The words of a text, without the marker syntax that export and import rewrite on purpose."""
    text = re.sub(r"(?m)^\s*:::.*$", "", text)  # block fences (an alias such as ::: note comes back as ::: written-note)
    text = re.sub(r"\[\[\w+:|\]\]", "", text)  # [[sc:..]] style markers (backticks come back as [[mono:..]])
    return re.findall(r"\w+", text.lower())


def test_no_words_are_lost_when_the_corpus_goes_through_docx_and_back(mw, corpus):
    for name, text in corpus.items():
        back, _notes = mw.import_docx_bytes(mw.create_manuscript_docx(text, {}))
        assert words(back) == words(CONTROL.sub("", text)), name


def test_control_characters_from_pdfs_do_not_break_the_exports(mw, corpus):
    text = "Page one" + chr(12) + "Page two" + chr(0) + " end\n\n# Chapter\n\nBell" + chr(7) + " ring"
    docx_text, _ = mw.import_docx_bytes(mw.create_manuscript_docx(text, {"title": "T" + chr(12)}))
    assert docx_text.startswith("Page onePage two end") and docx_text.endswith("Bell ring")
    z = zipfile.ZipFile(io.BytesIO(mw.create_manuscript_epub(text, {"title": "T" + chr(12)})))
    from lxml import etree

    for name in z.namelist():
        if name.endswith((".xhtml", ".opf", ".ncx")):
            etree.fromstring(z.read(name))
