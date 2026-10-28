"""DOCX export of Markdown lists (real Word lists, numbering that restarts) and pipe tables (real Word tables)."""
from __future__ import annotations

import io
import json
import zipfile

import pytest

NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def export(mw, text, **options):
    from docx import Document

    return Document(io.BytesIO(mw.create_manuscript_docx(text, {}, options or None)))


def body(doc):
    """The body in document order: ('p', style, text, numId) and ('table', rows) entries."""
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    out = []
    for child in doc.element.body.iterchildren():
        if child.tag == NS + "p":
            p = Paragraph(child, doc)
            num = child.find(f".//{NS}numId")
            out.append(("p", p.style.name, p.text, num.get(NS + "val") if num is not None else None))
        elif child.tag == NS + "tbl":
            out.append(("table", [[cell.text for cell in row.cells] for row in Table(child, doc).rows]))
    return out


def lines_of(doc):
    return [(entry[1], entry[2]) for entry in body(doc) if entry[0] == "p" and entry[2]]


def test_bullets_become_word_list_paragraphs_with_nesting(mw):
    doc = export(mw, "- one\n- two *it*\n  - nested a\n  - nested b\n    - deep\n- three")
    assert lines_of(doc) == [
        ("List Bullet", "one"), ("List Bullet", "two it"),
        ("List Bullet 2", "nested a"), ("List Bullet 2", "nested b"),
        ("List Bullet 3", "deep"), ("List Bullet", "three"),
    ]
    runs = {r.text: r for r in doc.paragraphs[1].runs}
    assert runs["it"].italic


def test_every_marker_style_and_a_tab_for_indentation_are_lists(mw):
    doc = export(mw, "* star\n+ plus\n1) paren\n\t- tabbed")
    assert [style for style, _ in lines_of(doc)] == ["List Bullet", "List Bullet", "List Number", "List Bullet 2"]


def test_numbered_lists_restart_and_keep_their_first_number(mw):
    doc = export(mw, "1. first\n2. second\n\nText in between.\n\n5. five\n6. six\n\n- bullets\n\n1. back to one")
    entries = [e for e in body(doc) if e[0] == "p" and e[2]]
    ids = {e[2]: e[3] for e in entries}
    assert ids["first"] == ids["second"] and ids["five"] == ids["six"]
    assert len({ids["first"], ids["five"], ids["back to one"]}) == 3
    xml = zipfile.ZipFile(io.BytesIO(doc_bytes(doc))).read("word/numbering.xml").decode()
    assert 'w:val="5"' in xml  # the list that starts at five


def doc_bytes(doc):
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def test_a_blank_line_between_items_does_not_end_the_list(mw):
    doc = export(mw, "1. a\n\n2. b\n\n3. c")
    ids = {e[3] for e in body(doc) if e[0] == "p" and e[1] == "List Number"}
    assert len(ids) == 1


def test_a_sublist_starts_again_under_every_item(mw):
    doc = export(mw, "1. one\n   1. sub a\n   2. sub b\n2. two\n   1. sub c")
    ids = {e[2]: e[3] for e in body(doc) if e[0] == "p" and e[2]}
    assert ids["sub a"] == ids["sub b"] and ids["sub c"] != ids["sub a"]
    assert ids["one"] == ids["two"]


def test_lists_can_be_switched_off(mw):
    doc = export(mw, "- a\n- b", lists=False)
    assert lines_of(doc) == [("Normal", "- a"), ("Normal", "- b")]


def test_list_like_lines_in_a_structure_block_stay_text(mw):
    doc = export(mw, "::: verse\n- not a list\n1. nor this\n:::\n\n- but this is")
    styles = dict((text, style) for style, text in lines_of(doc))
    assert styles["- not a list"] == "Manuscript Verse" and styles["1. nor this"] == "Manuscript Verse"
    assert styles["but this is"] == "List Bullet"


def test_scene_breaks_and_bold_text_are_not_lists(mw):
    doc = export(mw, "* * *\n\n**Bold start** and more\n\n---")
    assert all(not style.startswith("List") for style, _ in lines_of(doc))


def test_pipe_table_becomes_a_word_table_in_place(mw):
    source = "Before.\n\n| Name | Qty | Note |\n|---|:---:|---:|\n| Apple | 3 | **good** |\n| Pear | 10 | a \\| b |\n\nAfter."
    doc = export(mw, source)
    kinds = [entry[0] for entry in body(doc)]
    assert kinds.index("table") == 2 and body(doc)[kinds.index("table")][1] == [
        ["Name", "Qty", "Note"], ["Apple", "3", "good"], ["Pear", "10", "a | b"],
    ]
    assert [e[2] for e in body(doc) if e[0] == "p" and e[2]] == ["Before.", "After."]


def test_table_style_header_alignment_and_inline_markup(mw):
    doc = export(mw, "| A | B | C |\n|---|:---:|---:|\n| x | y | **z** |")
    table = doc.tables[0]
    assert table.style.name == "Table Grid"
    assert all(run.bold for cell in table.rows[0].cells for run in cell.paragraphs[0].runs)
    aligns = [str(cell.paragraphs[0].alignment) for cell in table.rows[1].cells]
    assert aligns == ["None", "CENTER (1)", "RIGHT (2)"]
    assert table.rows[1].cells[2].paragraphs[0].runs[0].bold
    assert "w:tblHeader" in zipfile.ZipFile(io.BytesIO(doc_bytes(doc))).read("word/document.xml").decode()


def test_ragged_rows_are_padded_and_br_is_a_line_break(mw):
    doc = export(mw, "| A | B |\n|---|---|\n| only one |\n| x | one<br>two | extra |")
    rows = [[cell.text for cell in row.cells] for row in doc.tables[0].rows]
    assert rows == [["A", "B", ""], ["only one", "", ""], ["x", "one\ntwo", "extra"]]


def test_table_like_lines_inside_a_structure_block_stay_text(mw):
    doc = export(mw, "::: verse\n| a | b |\n|---|---|\n:::")
    assert doc.tables == [] and "| a | b |" in [text for _, text in lines_of(doc)]


def test_two_tables_stay_apart(mw):
    kinds = [entry[0] for entry in body(export(mw, "| a | b |\n|---|---|\n| 1 | 2 |\n\n| c | d |\n|---|---|\n| 3 | 4 |"))]
    first = kinds.index("table")
    second = kinds.index("table", first + 1)
    assert "p" in kinds[first + 1:second]  # Word would merge two tables that touch


def test_export_endpoint_passes_the_lists_option(server):
    from docx import Document

    for flag, expected in ((True, "List Bullet"), (False, "Normal")):
        status, _h, data = server.request("POST", "/api/export_docx", body={"text": "- a", "metadata": {}, "options": {"lists": flag}})
        assert status == 200
        assert Document(io.BytesIO(data)).paragraphs[0].style.name == expected
